import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import json

# Crear un nuevo archivo Excel
workbook = Workbook()

# Hoja de Configuración
config_sheet = workbook.active
config_sheet.title = "Configuración"
config_sheet['A1'] = "Parámetro"
config_sheet['B1'] = "Valor"

config_data = [
    ["Título", "Encuesta de Satisfacción - Servicio de Restaurante"],
    ["Descripción", "Evalúe su experiencia en nuestro restaurante"],
    ["Preguntas por página", 5]
]

for row, data in enumerate(config_data, start=2):
    config_sheet[f'A{row}'] = data[0]
    config_sheet[f'B{row}'] = data[1]

# Hoja de Preguntas
questions_sheet = workbook.create_sheet("Preguntas")
questions_sheet['A1'] = "Pregunta"

questions = [
    "El tiempo de espera para ser atendido fue adecuado",
    "El personal fue amable y cortés",
    "El mesero conocía bien el menú",
    "El mesero fue atento durante toda la experiencia",
    "El servicio fue rápido y eficiente",
    "La comida estaba a la temperatura adecuada",
    "La presentación de los platos fue atractiva",
    "La calidad de los ingredientes fue excelente",
    "El sabor de la comida cumplió mis expectativas",
    "La porción servida fue apropiada",
    "El ambiente del restaurante era agradable",
    "El nivel de ruido era apropiado",
    "Las instalaciones estaban limpias",
    "La decoración era atractiva",
    "La iluminación era adecuada",
    "La relación calidad-precio es buena",
    "El proceso de pago fue eficiente",
    "Recomendaría este restaurante",
    "Volvería a visitar este restaurante",
    "Mi experiencia general fue satisfactoria"
]

for row, question in enumerate(questions, start=2):
    questions_sheet[f'A{row}'] = question

# Hoja de Categorías
categories_sheet = workbook.create_sheet("Categorías")
categories_sheet['A1'] = "Categoría"
categories_sheet['B1'] = "Índices de preguntas"

categories = [
    ["Servicio", [0, 1, 2, 3, 4]],
    ["Comida", [5, 6, 7, 8, 9]],
    ["Ambiente", [10, 11, 12, 13, 14]],
    ["Satisfacción General", [15, 16, 17, 18, 19]]
]

for row, (category, indices) in enumerate(categories, start=2):
    categories_sheet[f'A{row}'] = category
    categories_sheet[f'B{row}'] = json.dumps(indices)

# Dar formato al Excel
for sheet in workbook:
    # Formato de encabezados
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    # Ajustar ancho de columnas
    for column in sheet.columns:
        max_length = 0
        column = list(column)
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        sheet.column_dimensions[column[0].column_letter].width = adjusted_width

# Guardar el archivo
workbook.save('restaurant_survey.xlsx')